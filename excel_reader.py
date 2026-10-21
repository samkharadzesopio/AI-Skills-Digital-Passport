
from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"

excel_files = list(DATA_DIR.glob("*.xlsx"))

if not excel_files:
    raise FileNotFoundError("No Excel file found in data folder.")

workbook = load_workbook(
    excel_files[0],
    read_only=True,
    data_only=True
)

sheet = workbook["Passport Skill Records"]


def get_student(student_id):
    student = None
    seen_records = set()

    # Row 4 contains headers, data starts at row 5
    for row in sheet.iter_rows(
        min_row=5,
        values_only=True
    ):
        # Column A - Student ID
        if row[0] != student_id:
            continue

        # Create student only once
        if student is None:
            student = {
                "student_id": row[0],
                "first_name": row[1],
                "last_name": row[2],
                "email": row[3],
                "passport_id": row[6],
                "first_issued": row[7],
                "skills": []
            }

        # Column J - Record ID (unique award event)
        record_id = row[9]

        if record_id in seen_records:
            continue

        seen_records.add(record_id)

        # Add every awarded skill
        student["skills"].append({
            "record_id": record_id,
            "skill_id": row[10],
            "name": row[11],
            "category": row[12],
            "badge": row[13],
            "description": row[14],
            "earned_through": "",
            "date_earned": row[15]
        })

    return student

if __name__ == "__main__":
    student = get_student("BTU-T-0001")

    if student:
        print("Student:", student["first_name"], student["last_name"])
        print("Passport ID:", student["passport_id"])
        print("Number of skills:", len(student["skills"]))

        for skill in student["skills"]:
            print("-", skill["name"])
    else:
        print("Student not found.")


