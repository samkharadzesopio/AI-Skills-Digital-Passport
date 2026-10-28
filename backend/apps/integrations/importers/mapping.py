"""Bounded, deterministic schema mapping. No database writes or approximate matching."""
import hashlib
import json
import math
import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .excel import MAX_FILE_SIZE, MAX_ROWS
from .validators import ImportValidationError, REQUIRED_FIELDS, validate_completion_record

MAX_COLUMNS = 26
HEADER_SCAN = 10
SKILL_FIELDS = {'external_student_id', 'source_event_id', 'skill_code', 'earned_on'}
OPTIONAL_FIELDS = {'approval_status', 'approval_reference', 'first_name', 'last_name', 'email',
                   'student_status', 'skill_name', 'category', 'stamp_type'}
OUTPUT_HEADERS = {'passport sync status', 'synced passport version', 'synced at utc',
                  'sync error notes', 'passport id existing', 'first issued existing'}


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, separators=(',', ':'),
                                     ensure_ascii=True, allow_nan=False).encode()).hexdigest()


def normalize(value):
    return re.sub(r'[\s_./()\-]+', ' ', unicodedata.normalize('NFKC', str(value or '')).strip().casefold()).strip()


def validate_definition(kind, definition):
    required = set(REQUIRED_FIELDS) if kind == 'COMPLETION' else SKILL_FIELDS
    if kind not in ('COMPLETION', 'SKILL_AWARD') or not isinstance(definition, dict):
        raise ImportValidationError('Invalid profile kind or definition.')
    if set(definition) - {'columns', 'value_maps', 'date_formats', 'allowed_skill_codes', 'eligible_student_statuses', 'worksheet', 'google_source', 'evidence_namespace'}:
        raise ImportValidationError('Unsupported profile option.')
    columns = definition.get('columns')
    if not isinstance(columns, dict) or not required.issubset(columns) or set(columns) - (required | OPTIONAL_FIELDS):
        raise ImportValidationError('Profile must map the required canonical fields.')
    seen = {}
    for field, aliases in columns.items():
        if not isinstance(aliases, list) or not 1 <= len(aliases) <= 10:
            raise ImportValidationError('Each field needs 1–10 explicit header aliases.')
        for alias in aliases:
            if not isinstance(alias, str) or not alias.strip() or len(alias) > 150:
                raise ImportValidationError('Invalid header alias.')
            key = normalize(alias)
            if key in OUTPUT_HEADERS or (key in seen and seen[key] != field):
                raise ImportValidationError('Ambiguous alias or system output used as evidence.')
            seen[key] = field
    maps = definition.get('value_maps', {})
    if not isinstance(maps, dict) or set(maps) - {'approval_status', 'completion_status', 'passed', 'skill_code', 'student_status', 'stamp_type', 'category'}:
        raise ImportValidationError('Invalid controlled value maps.')
    for field, values in maps.items():
        if not isinstance(values, dict) or len(values) > 500:
            raise ImportValidationError('Invalid value map.')
        keys = set()
        for key, target in values.items():
            if not isinstance(key, str) or not key.strip() or len(key) > 150 or not isinstance(target, (str, bool)):
                raise ImportValidationError('Invalid value mapping.')
            if isinstance(target, str) and (not target.strip() or len(target) > 255):
                raise ImportValidationError('Invalid mapped value.')
            canonical = key.strip() if field == 'skill_code' else normalize(key)
            if canonical in keys or (field == 'approval_status' and (
                target not in ('CONFIRMED', 'PENDING', 'REJECTED') or
                canonical in ('pending', 'rejected', 'synced', 'not synced') and target == 'CONFIRMED'
            )):
                raise ImportValidationError('Ambiguous or unsafe approval mapping.')
            keys.add(canonical)
    formats = definition.get('date_formats', ['%Y-%m-%d'])
    if not isinstance(formats, list) or not 1 <= len(formats) <= 5 or any(
        f not in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%d.%m.%Y') for f in formats
    ):
        raise ImportValidationError('Unsupported date format.')
    if kind == 'SKILL_AWARD':
        if not isinstance(definition.get('evidence_namespace'), str) or not re.fullmatch(r'[a-z0-9_-]{1,80}', definition['evidence_namespace']):
            raise ImportValidationError('Skill awards require a stable institutional evidence namespace across exports.')
        codes = definition.get('allowed_skill_codes')
        if not isinstance(codes, list) or not codes or len(codes) > 500 or any(
            not isinstance(c, str) or not c.strip() or c != c.strip() or len(c) > 100 for c in codes
        ) or len(set(codes)) != len(codes):
            raise ImportValidationError('Skill awards require an explicit authorized catalog allowlist.')
        statuses = definition.get('eligible_student_statuses')
        if not isinstance(statuses, list) or not statuses or len(statuses) > 20 or any(
            not isinstance(s, str) or not s.strip() or len(s) > 50 for s in statuses
        ):
            raise ImportValidationError('Skill awards require explicit eligible Student statuses.')
    sheet = definition.get('worksheet', 'Completions')
    if not isinstance(sheet, str) or not 1 <= len(sheet) <= 100:
        raise ImportValidationError('Invalid worksheet name.')
    if 'google_source' in definition and not re.fullmatch(r'[A-Za-z0-9_-]{1,80}', str(definition['google_source'])):
        raise ImportValidationError('Invalid server-side Google source key.')
    if len(json.dumps(definition).encode()) > 100_000:
        raise ImportValidationError('Profile exceeds processing limits.')
    return definition


def excel_grid(file, *, worksheet):
    content = file.read(MAX_FILE_SIZE + 1)
    if not content or len(content) > MAX_FILE_SIZE:
        raise ImportValidationError('Excel file must be nonempty and at most 5 MiB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
                raise ImportValidationError('Workbook expansion exceeds processing limits.')
            if any('externallinks/' in e.filename.lower() or 'vbaproject' in e.filename.lower()
                   or '/embeddings/' in e.filename.lower() for e in entries):
                raise ImportValidationError('External links, embedded objects and macros are not supported.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ImportValidationError:
        raise
    except Exception:
        raise ImportValidationError('Invalid Excel workbook.') from None
    try:
        if worksheet not in book.sheetnames:
            raise ImportValidationError('The configured worksheet is missing.')
        sheet = book[worksheet]
        # Do not trust missing or deliberately incorrect XML dimension metadata.
        sheet.reset_dimensions()
        rows = []
        for row in sheet.iter_rows(max_col=MAX_COLUMNS + 1, max_row=MAX_ROWS + HEADER_SCAN + 1):
            if row[-1].value is not None:
                raise ImportValidationError('Workbook exceeds 26-column limit.')
            values = [c.value for c in row[:-1]]
            # Formula objects are never evaluated; mapped formula strings fail row validation.
            if any(c.data_type == 'f' and not isinstance(c.value, str) for c in row[:-1]):
                raise ImportValidationError('Unsupported spreadsheet formula.')
            rows.append(values)
        rows = bounded_grid(rows)
        return rows, hashlib.sha256(content).hexdigest()
    except ImportValidationError:
        raise
    except Exception:
        raise ImportValidationError('Invalid Excel workbook.') from None
    finally:
        book.close()


def bounded_grid(rows):
    if not isinstance(rows, list) or len(rows) > MAX_ROWS + HEADER_SCAN + 1:
        raise ImportValidationError('Source exceeds row limit.')
    output = []
    for row in rows:
        if not isinstance(row, (list, tuple)) or len(row) > MAX_COLUMNS:
            raise ImportValidationError('Source exceeds 26-column limit.')
        normalized = []
        for value in row:
            if isinstance(value, datetime):
                value = value.date()
            if isinstance(value, date):
                value = value.isoformat()
            if value is not None and not isinstance(value, (str, int, float, bool)):
                raise ImportValidationError('Unsupported cell type.')
            if isinstance(value, float) and not math.isfinite(value):
                raise ImportValidationError('Nonfinite spreadsheet number.')
            normalized.append(value)
        if sum(len(str(v)) for v in normalized if v is not None) > 16384:
            raise ImportValidationError('Source row exceeds processing limit.')
        # Canonical padding makes Sheets/Excel snapshot comparisons deterministic.
        output.append(normalized + [None] * (MAX_COLUMNS - len(normalized)))
    while output and not any(v is not None and str(v).strip() for v in output[-1]):
        output.pop()
    return output


def normalized_date(value, formats):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if not isinstance(value, str):
        raise ImportValidationError('INVALID_DATE')
    candidates = set()
    for fmt in formats:
        try:
            candidates.add(datetime.strptime(value.strip(), fmt).date())
        except ValueError:
            continue
    if len(candidates) != 1 or next(iter(candidates)) > date.today():
        raise ImportValidationError('INVALID_OR_AMBIGUOUS_DATE')
    return candidates.pop().isoformat()


def map_grid(rows, *, kind, definition):
    validate_definition(kind, definition)
    rows = bounded_grid(rows)
    required = set(REQUIRED_FIELDS) if kind == 'COMPLETION' else SKILL_FIELDS
    aliases = {normalize(alias): field for field, items in definition['columns'].items() for alias in items}
    candidates = []
    for number, row in enumerate(rows[:HEADER_SCAN], 1):
        columns, ambiguous = {}, False
        for index, cell in enumerate(row):
            field = aliases.get(normalize(cell))
            if field:
                if field in columns:
                    ambiguous = True
                columns[field] = index
        if required.issubset(columns):
            if ambiguous:
                raise ImportValidationError('Ambiguous header mapping.')
            candidates.append((number, columns))
    if len(candidates) != 1:
        raise ImportValidationError('Expected one unambiguous header row in the first 10 rows.')
    header, columns = candidates[0]
    if len(rows) - header > MAX_ROWS:
        raise ImportValidationError('Source exceeds 5000 data rows.')
    records, issues, invalid_event_ids = [], [], set()
    value_maps = {field: {(k.strip() if field == 'skill_code' else normalize(k)): v for k, v in mapping.items()}
                  for field, mapping in definition.get('value_maps', {}).items()}
    for number, row in enumerate(rows[header:], header + 1):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        record = {}
        try:
            record = {field: row[index].strip() if isinstance(row[index], str) else row[index]
                      for field, index in columns.items()}
            if any(isinstance(v, str) and v.startswith(('=', '+', '@')) for v in record.values()):
                raise ImportValidationError('FORMULA_NOT_ALLOWED')
            for field, mapping in value_maps.items():
                if field in record:
                    lookup = str(record[field]).strip() if field == 'skill_code' else normalize(record[field])
                    record[field] = mapping.get(lookup, record[field])
            for field in required - {'grade', 'passed', 'completion_date', 'earned_on'}:
                v = record.get(field)
                if not isinstance(v, str) or not v or len(v) > 120 or any(ord(c) < 32 for c in v):
                    raise ImportValidationError('MISSING_OR_INVALID_IDENTIFIER')
            field = 'completion_date' if kind == 'COMPLETION' else 'earned_on'
            record[field] = normalized_date(record[field], definition.get('date_formats', ['%Y-%m-%d']))
            if kind == 'COMPLETION':
                try:
                    if isinstance(record['grade'], str):
                        record['grade'] = float(record['grade'])
                except ValueError:
                    raise ImportValidationError('INVALID_GRADE') from None
                if isinstance(record['passed'], str) and record['passed'].casefold() in ('true', 'false'):
                    record['passed'] = record['passed'].casefold() == 'true'
                extras = {k: v for k, v in record.items() if k in OPTIONAL_FIELDS}
                record = {**validate_completion_record(record), **extras}
                record['grade'] = float(record['grade'])
            elif record['skill_code'] not in definition['allowed_skill_codes']:
                raise ImportValidationError('SKILL_NOT_AUTHORIZED')
            records.append({'row': number, 'record': record})
        except (ImportValidationError, ValueError) as exc:
            identity = record.get('source_event_id')
            if isinstance(identity, str) and identity.strip():
                invalid_event_ids.add(identity.strip())
            # Never echo uploaded values or library exception text.
            code = str(exc) if re.fullmatch('[A-Z_]{1,64}', str(exc)) else 'INVALID_COMPLETION_ROW'
            issues.append({'row': number, 'status': 'INVALID', 'code': code})
    return {'header_row': header, 'columns': columns, 'records': records, 'issues': issues, 'invalid_event_ids': invalid_event_ids,
            'total_rows': len(records) + len(issues)}
