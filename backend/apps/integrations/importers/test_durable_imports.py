from datetime import timedelta
from django.utils import timezone
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO, StringIO
from threading import Barrier
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from django.core.management import call_command
from django.core.exceptions import ImproperlyConfigured
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections
from django.test import TransactionTestCase, SimpleTestCase, override_settings
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.credentials.models import Credential
from apps.credentials.test_support import create_test_issuer
from apps.skills.models import Skill, VerifiedAchievement
from apps.students.models import Student
from apps.integrations.models import ImportRun, ClassroomSyncEvent
from .durable import process_import, reconcile_run
from .excel import read_excel_completions
from .validators import ImportValidationError, REQUIRED_FIELDS

ROW = {'external_student_id': 'IMPORT-SYNTHETIC-001', 'source_event_id': 'IMPORT-EVENT-001',
       'course_id': 'AI-101', 'completion_status': 'completed', 'grade': 90, 'passed': True,
       'completion_date': '2026-01-01'}


@override_settings(BTU_CLASSROOM_ISSUER_REFERENCE='BTU-UNIVERSITY', BTU_AUTHORIZED_COMPLETION_SOURCES=['BTU_EXCEL', 'BTU_GOOGLE_SHEETS'])
class DurableImportTests(TransactionTestCase):
    def setUp(self):
        self.student = Student.objects.create(external_student_id=ROW['external_student_id'], first_name='Synthetic', last_name='Student',
                                             email='import@btu.edu.ge', program='AI', education_level='Bachelor', student_status='Active')
        create_test_issuer('Synthetic BTU')
        for code in ['AI_LITERACY', 'AI_FUNDAMENTALS']:
            Skill.objects.create(skill_code=code, skill_name=code, category='AI Skill', description='Synthetic', stamp_type='FOCUSED', badge_asset_id='test')

    def test_real_import_issues_and_repeated_rows_are_idempotent(self):
        result = process_import([ROW], source='BTU_EXCEL')
        self.assertEqual((result['issued'], result['awarded']), (1, 2))
        repeat = process_import([{**ROW, 'grade': 90.0}], source='BTU_EXCEL')
        self.assertEqual(repeat['skipped'], 1)
        self.assertEqual(Credential.objects.count(), 1)
        self.assertEqual(ImportRun.objects.count(), 2)
        self.assertEqual(ClassroomSyncEvent.objects.count(), 1)

    def test_partial_failure_replay_resumes_issuance(self):
        with patch('apps.integrations.classroom_automation.issue_credential', side_effect=RuntimeError('private database details')):
            result = process_import([ROW], source='BTU_GOOGLE_SHEETS')
        self.assertEqual(result['failed'], 1)
        self.assertEqual(VerifiedAchievement.objects.count(), 2)
        self.assertEqual(ImportRun.objects.get().state, 'FAILED')
        event = ClassroomSyncEvent.objects.get()
        from apps.integrations.classroom_automation import process_event
        self.assertEqual(process_event(event.pk, force=True), 'issued')
        run = ImportRun.objects.get()
        reconcile_run(run)
        self.assertEqual(run.state, 'COMPLETED')
        self.assertEqual(run.summary['current_issued'], 1)
        self.assertEqual(VerifiedAchievement.objects.count(), 2)

    @override_settings(BTU_AUTHORIZED_COMPLETION_SOURCES=[])
    def test_default_untrusted_spreadsheet_cannot_award(self):
        with self.assertRaises(ValueError):
            process_import([ROW], source='BTU_EXCEL')
        self.assertFalse(VerifiedAchievement.objects.exists())
        self.assertFalse(Credential.objects.exists())

    def test_conflicting_payload_is_reported_without_private_exception(self):
        process_import([ROW], source='BTU_EXCEL')
        result = process_import([{**ROW, 'grade': 95}], source='BTU_EXCEL')
        self.assertEqual(result['failed'], 1)
        self.assertEqual(result['errors'][0]['type'], 'EventConflict')
        self.assertEqual(Credential.objects.count(), 1)

    def test_invalid_row_is_tracked_with_valid_partial_batch(self):
        result = process_import([ROW, {**ROW, 'source_event_id': 'BAD-ROW', 'grade': 0}], source='BTU_EXCEL')
        self.assertEqual((result['issued'], result['failed']), (1, 1))
        self.assertEqual(ImportRun.objects.get().state, 'PARTIAL')

    def test_concurrent_imports_share_one_credential(self):
        barrier = Barrier(2)
        def run(_):
            close_old_connections()
            try:
                barrier.wait(timeout=10)
                return process_import([ROW], source='BTU_EXCEL')
            finally:
                close_old_connections()
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(run, range(2)))
        self.assertEqual(sum(r['issued'] for r in results), 1)
        self.assertEqual(sum(r['skipped'] for r in results), 1)
        self.assertEqual(Credential.objects.count(), 1)

    def test_reconciliation_detects_interrupted_fetch(self):
        run = ImportRun.objects.create(source='BTU_EXCEL')
        ImportRun.objects.filter(pk=run.pk).update(created_at=timezone.now() - timedelta(minutes=10))
        call_command('reconcile_import_runs', stdout=StringIO())
        run.refresh_from_db()
        self.assertEqual(run.state, 'FAILED')
        self.assertEqual(run.summary['error'], 'INTERRUPTED_BEFORE_EVENT_RECEIPT')

    def test_real_excel_api_history_and_permission(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(list(REQUIRED_FIELDS))
        sheet.append([ROW[k] for k in REQUIRED_FIELDS])
        sheet.append([None, None, 'in_progress'])
        file = BytesIO()
        workbook.save(file)
        client = APIClient()
        admin = User.objects.create_admin_user(email='import.admin@btu.edu.ge', password='Synthetic-test-only-2026!', permission_codenames={'import_completions'})
        client.force_authenticate(admin)
        response = client.post('/api/integrations/import/excel/', {'file': SimpleUploadedFile('records.xlsx', file.getvalue())}, format='multipart')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['result']['issued'], 1)
        run = ImportRun.objects.get()
        self.assertEqual(run.state, 'PARTIAL')
        self.assertEqual(len(run.validation_errors), 1)
        self.assertEqual(client.get('/api/integrations/imports/').status_code, 200)
        event = ClassroomSyncEvent.objects.get()
        self.assertEqual(client.post(f'/api/integrations/imports/events/{event.pk}/retry/', {}, format='json').json()['result'], 'skipped')
        # Import permission cannot operate the separate Classroom retry route.
        self.assertEqual(client.post(f'/api/integrations/classroom/events/{event.pk}/retry/', {}, format='json').status_code, 403)

    def test_failed_workbook_leaves_durable_history(self):
        client = APIClient()
        admin = User.objects.create_admin_user(email='badfile.admin@btu.edu.ge', password='Synthetic-test-only-2026!', permission_codenames={'import_completions'})
        client.force_authenticate(admin)
        response = client.post('/api/integrations/import/excel/', {'file': SimpleUploadedFile('records.xlsx', b'not a workbook')}, format='multipart')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(ImportRun.objects.get().state, 'FAILED')
        self.assertFalse(Credential.objects.exists())

    def test_legacy_integer_grade_achievement_is_reused(self):
        from datetime import date
        from apps.skills.achievement_services import record_verified_achievements
        awards, created = record_verified_achievements(
            source='BTU_EXCEL', source_event_id=ROW['source_event_id'],
            external_student_id=ROW['external_student_id'],
            approved_skill_codes=['AI_LITERACY', 'AI_FUNDAMENTALS'],
            earned_on=date.fromisoformat(ROW['completion_date']), verification_status='VERIFIED',
            course_id=ROW['course_id'], evidence_reference=f"BTU_EXCEL:{ROW['source_event_id']}",
            evidence={'grade': 90, 'passed': True, 'completion_status': 'completed'},
        )
        self.assertTrue(created)
        before = {row.pk for row in awards}
        result = process_import([ROW], source='BTU_EXCEL')
        self.assertEqual(result['issued'], 1)
        self.assertEqual(set(Credential.objects.get().achievements.values_list('pk', flat=True)), before)


class WorkbookLimitsTests(SimpleTestCase):
    def test_zip_expansion_and_column_limits(self):
        file = BytesIO()
        with ZipFile(file, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('large.xml', b' ' * (32 * 1024 * 1024 + 1))
        file.seek(0)
        with self.assertRaises(ImportValidationError):
            read_excel_completions(file)
        book = Workbook()
        book.active.cell(1, 27, 'too wide')
        file = BytesIO()
        book.save(file)
        file.seek(0)
        with self.assertRaises(ImportValidationError):
            read_excel_completions(file)


class PreviewAuditRegressionTests(TransactionTestCase):
    def test_authorized_preview_does_not_write_audit_records(self):
        from apps.accounts.models import SecurityAudit
        client = APIClient()
        admin = User.objects.create_admin_user(email='preview.admin@btu.edu.ge', password='Synthetic-test-only-2026!', permission_codenames={'sync_classroom'})
        client.force_authenticate(admin)
        with patch('apps.integrations.classroom_automation.ClassroomService.get_completions', return_value=[]):
            self.assertEqual(client.post('/api/integrations/classroom/sync/', {'dry_run': True}, format='json').status_code, 200)
        self.assertFalse(SecurityAudit.objects.exists())
        self.assertFalse(ClassroomSyncEvent.objects.exists())
        self.assertFalse(Credential.objects.exists())

    def test_denied_request_does_not_parse_unsupported_body_for_audit(self):
        client = APIClient()
        response = client.generic('POST', '/api/integrations/classroom/sync/', b'not-json', content_type='application/x-unsupported')
        self.assertEqual(response.status_code, 403)
        from apps.accounts.models import SecurityAudit
        self.assertEqual(SecurityAudit.objects.get().status_code, 403)

    @override_settings(BTU_AUTHORIZED_COMPLETION_SOURCES=['BTU_EXCEL'])
    def test_classroom_preview_cannot_claim_spreadsheet_provenance(self):
        from apps.integrations.classroom_automation import ClassroomCredentialAutomation
        from unittest.mock import Mock
        service = Mock()
        service.get_completions.return_value = [{'source': 'BTU_EXCEL'}]
        result = ClassroomCredentialAutomation(service).synchronize(dry_run=True)
        self.assertEqual(result['failed'], 1)
        self.assertEqual(result['would_issue'], 0)
