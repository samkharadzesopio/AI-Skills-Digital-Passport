
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook

from .validators import (
    REQUIRED_FIELDS,
    validate_completion_record,
    ImportValidationError,
)


MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 5000


def read_excel_completions(file):
    """
    Read an Excel workbook and validate
    completion records.

    Returns:
        valid_records, errors
    """

    content = file.read(MAX_FILE_SIZE + 1)

    if len(content) > MAX_FILE_SIZE:
        raise ImportValidationError(
            "Excel file exceeds 5 MB limit"
        )

    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 32 * 1024 * 1024:
                raise ImportValidationError('Workbook expansion exceeds processing limits')
    except BadZipFile:
        raise ImportValidationError('Invalid Excel workbook') from None

    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ImportValidationError(
            "Invalid Excel workbook"
        ) from exc

    valid_records = []
    errors = []

    try:
        sheet = workbook.active
        if (sheet.max_column or 0) > 26:
            raise ImportValidationError('Workbook exceeds 26-column limit')
        rows = sheet.iter_rows(values_only=True, max_col=min(sheet.max_column or 26, 26))

        headers = next(rows, None)

        if not headers:
            raise ImportValidationError(
                "Excel sheet is empty"
            )

        headers = [
            str(value).strip()
            if value is not None else ""
            for value in headers
        ]

        while headers and not headers[-1]:
            headers.pop()
        if not set(REQUIRED_FIELDS).issubset(headers):
            raise ImportValidationError('Missing required completion columns')
        if len(set(headers)) != len(headers):
            raise ImportValidationError(
                "Duplicate column names"
            )

        for row_number, values in enumerate(
            rows, start=2
        ):
            if row_number > MAX_ROWS + 1:
                raise ImportValidationError(
                    "Excel file exceeds row limit"
                )

            if not any(
                value is not None for value in values
            ):
                continue

            if sum(len(str(value)) for value in values if value is not None) > 16384:
                raise ImportValidationError('Workbook row exceeds processing limit')
            record = dict(zip(headers, values))

            try:
                validated = validate_completion_record(
                    record
                )
                valid_records.append(validated)
            except (ImportValidationError, ValueError) as exc:
                errors.append({
                    "row": row_number,
                    "error": str(exc),
                })

    finally:
        workbook.close()

    return valid_records, errors
