
from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"


def format_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")

    return str(value) if value is not None else ""


def get_student(student_id):
    excel_files = sorted(DATA_DIR.glob("*.xlsx"))

    if not excel_files:
        raise FileNotFoundError(
            "No Excel file found in data folder."
        )

    workbook = load_workbook(
        excel_files[0],
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook["Passport Skill Records"]

        student = None
        seen_records = set()

        # Row 4 contains headers; records start at row 5.
        for row in sheet.iter_rows(
            min_row=5,
            values_only=True
        ):
            if row[0] != student_id:
                continue

            if student is None:
                student = {
                    "student_id": row[0],
                    "first_name": row[1],
                    "last_name": row[2],
                    "email": row[3],
                    "passport_id": row[6],
                    "first_issued": format_date(row[7]),
                    "skills": []
                }

            record_id = row[9]

            if not record_id:
                continue

            if record_id in seen_records:
                continue

            seen_records.add(record_id)

            student["skills"].append({
                "record_id": record_id,
                "skill_id": row[10],
                "name": row[11],
                "category": row[12],
                "badge": row[13],
                "description": row[14] or "",
                "earned_through": "",
                "date_earned": format_date(row[15])
            })

        return student

    finally:
        workbook.close()


if __name__ == "__main__":
    student_id = input("Enter Student ID: ").strip()

    student = get_student(student_id)

    if student:
        print("\nStudent:", student["first_name"], student["last_name"])
        print("Passport ID:", student["passport_id"])
        print("First issued:", student["first_issued"])
        print("Number of skills:", len(student["skills"]))

        for skill in student["skills"]:
            print(
                "-",
                skill["name"],
                "|",
                skill["category"],
                "|",
                skill["date_earned"]
            )
    else:
        print("Student not found.")
